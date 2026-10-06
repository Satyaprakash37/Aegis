"""Report generation engine for AEGIS platform.

Generates professional Executive PDF, Detailed Technical PDF, and Compliance Excel reports
using ReportLab and openpyxl.
"""

import os
from datetime import datetime
from typing import Tuple, List, Optional
import xml.sax.saxutils as saxutils

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models.asset import Asset
from app.models.vulnerability import Vulnerability
from app.models.scan import Scan
from app.models.user import User
from app.models.report import ReportType, ReportFormat
from app.services.risk.engine import get_risk_tier

# ReportLab Imports
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    PageBreak,
    KeepTogether,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas

# openpyxl Imports
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


REPORTS_DIR = os.getenv("REPORTS_DIR", "/app/reports")


def _get_reports_dir() -> str:
    """Ensure destination reports directory exists and return absolute path."""
    dest_dir = REPORTS_DIR
    if not os.path.exists(dest_dir):
        # Fallback to local project directory if not in Docker container
        local_fallback = os.path.join(
            os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(__file__)))),
            "backend",
            "reports",
        )
        os.makedirs(local_fallback, exist_ok=True)
        dest_dir = local_fallback
    else:
        os.makedirs(dest_dir, exist_ok=True)
    return dest_dir


def _clean_text(text: Optional[str]) -> str:
    """Safely escape XML characters for ReportLab Paragraphs."""
    if not text:
        return ""
    return saxutils.escape(str(text))


class NumberedCanvas(canvas.Canvas):
    """Two-pass canvas for dynamic 'Page X of Y' pagination and running headers/footers."""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count: int):
        # Suppress running header and footer on cover page
        if self._pageNumber == 1:
            return

        self.saveState()

        # Running Header
        self.setFont("Helvetica-Bold", 8)
        self.setFillColor(colors.HexColor("#06b6d4"))
        self.drawString(54, 752, "AEGIS SEC-OPS")
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748b"))
        self.drawString(130, 752, "| Continuous Vulnerability & Risk Management Platform")

        self.setStrokeColor(colors.HexColor("#334155"))
        self.setLineWidth(0.5)
        self.line(54, 744, 558, 744)

        # Running Footer
        self.line(54, 45, 558, 45)
        self.setFont("Helvetica-Bold", 7)
        self.setFillColor(colors.HexColor("#e11d48"))
        self.drawString(54, 34, "CONFIDENTIAL // RESTRICTED DISTRIBUTION")
        self.setFont("Helvetica", 7)
        self.setFillColor(colors.HexColor("#64748b"))
        self.drawString(240, 34, "AEGIS Security Audit Intelligence")

        page_str = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(558, 34, page_str)

        self.restoreState()


async def generate_report(
    report_type: ReportType,
    format: ReportFormat,
    user: User,
    db: AsyncSession,
) -> Tuple[str, str, int]:
    """Generate requested security report and persist to storage volume.

    Returns:
        (file_path, filename, file_size_bytes)
    """
    reports_dir = _get_reports_dir()
    timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    date_slug = datetime.utcnow().strftime("%Y-%m-%d")

    # Fetch data
    assets_stmt = select(Asset).order_by(Asset.criticality.desc(), Asset.name.asc())
    assets: List[Asset] = (await db.execute(assets_stmt)).scalars().all()

    vulns_stmt = (
        select(Vulnerability)
        .options(selectinload(Vulnerability.asset))
        .order_by(Vulnerability.risk_score.desc(), Vulnerability.cvss_score.desc())
    )
    vulns: List[Vulnerability] = (await db.execute(vulns_stmt)).scalars().all()

    scans_stmt = (
        select(Scan)
        .options(selectinload(Scan.asset))
        .order_by(Scan.started_at.desc())
        .limit(25)
    )
    scans: List[Scan] = (await db.execute(scans_stmt)).scalars().all()

    if format == ReportFormat.pdf:
        filename = f"aegis_{report_type.value}_{timestamp}.pdf"
        file_path = os.path.join(reports_dir, filename)

        if report_type == ReportType.executive:
            _build_executive_pdf(file_path, user, assets, vulns, date_slug)
        else:
            _build_detailed_pdf(file_path, user, assets, vulns, scans, date_slug)

    elif format == ReportFormat.excel:
        filename = f"aegis_compliance_{timestamp}.xlsx"
        file_path = os.path.join(reports_dir, filename)
        _build_compliance_excel(file_path, user, assets, vulns, date_slug)
    else:
        raise ValueError(f"Unsupported export format: {format}")

    file_size = os.path.getsize(file_path)
    return file_path, filename, file_size


# ==============================================================================
# PDF Generation Helpers
# ==============================================================================

def _get_pdf_stylesheet():
    """Build consistent styling dictionary for PDF reports."""
    base_styles = getSampleStyleSheet()

    custom_styles = {
        "CoverTitle": ParagraphStyle(
            "CoverTitle",
            parent=base_styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=24,
            leading=28,
            textColor=colors.HexColor("#0f172a"),
            alignment=0,
            spaceAfter=8,
        ),
        "CoverSubtitle": ParagraphStyle(
            "CoverSubtitle",
            parent=base_styles["Normal"],
            fontName="Helvetica",
            fontSize=12,
            leading=16,
            textColor=colors.HexColor("#475569"),
            alignment=0,
            spaceAfter=20,
        ),
        "SectionHeader": ParagraphStyle(
            "SectionHeader",
            parent=base_styles["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=13,
            leading=17,
            textColor=colors.HexColor("#0f172a"),
            spaceBefore=16,
            spaceAfter=8,
            keepWithNext=True,
        ),
        "SubsectionHeader": ParagraphStyle(
            "SubsectionHeader",
            parent=base_styles["Heading3"],
            fontName="Helvetica-Bold",
            fontSize=10,
            leading=14,
            textColor=colors.HexColor("#1e293b"),
            spaceBefore=10,
            spaceAfter=4,
            keepWithNext=True,
        ),
        "BodyDark": ParagraphStyle(
            "BodyDark",
            parent=base_styles["Normal"],
            fontName="Helvetica",
            fontSize=9,
            leading=13,
            textColor=colors.HexColor("#334155"),
            spaceAfter=6,
        ),
        "BodyDarkBold": ParagraphStyle(
            "BodyDarkBold",
            parent=base_styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=9,
            leading=13,
            textColor=colors.HexColor("#0f172a"),
        ),
        "TableHead": ParagraphStyle(
            "TableHead",
            parent=base_styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=8,
            leading=10,
            textColor=colors.white,
            alignment=0,
        ),
        "TableCell": ParagraphStyle(
            "TableCell",
            parent=base_styles["Normal"],
            fontName="Helvetica",
            fontSize=8,
            leading=11,
            textColor=colors.HexColor("#1e293b"),
        ),
        "TableCellBold": ParagraphStyle(
            "TableCellBold",
            parent=base_styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=8,
            leading=11,
            textColor=colors.HexColor("#0f172a"),
        ),
        "TableCellCode": ParagraphStyle(
            "TableCellCode",
            parent=base_styles["Normal"],
            fontName="Courier-Bold",
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#0369a1"),
        ),
        "MetaLabel": ParagraphStyle(
            "MetaLabel",
            parent=base_styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=8,
            leading=11,
            textColor=colors.HexColor("#64748b"),
        ),
        "MetaValue": ParagraphStyle(
            "MetaValue",
            parent=base_styles["Normal"],
            fontName="Helvetica",
            fontSize=8,
            leading=11,
            textColor=colors.HexColor("#0f172a"),
        ),
    }

    base_styles.byName.update(custom_styles)
    return base_styles


def _severity_color(severity_str: str) -> colors.HexColor:
    s = (severity_str or "").lower()
    if s == "critical":
        return colors.HexColor("#e11d48")
    elif s == "high":
        return colors.HexColor("#ea580c")
    elif s == "medium":
        return colors.HexColor("#ca8a04")
    elif s == "low":
        return colors.HexColor("#0284c7")
    return colors.HexColor("#64748b")


def _build_cover_page(story: list, styles, title: str, subtitle: str, user: User, total_assets: int, date_slug: str):
    """Assemble standard AEGIS branded cover page."""
    story.append(Spacer(1, 40))

    # Top brand badge
    badge_data = [[
        Paragraph("<font color='#0891b2'><b>AEGIS SECURITY COMMAND</b></font> &nbsp;|&nbsp; <font color='#64748b'>DEFENSE INTELLIGENCE</font>", styles["TableCellBold"])
    ]]
    badge_table = Table(badge_data, colWidths=[504])
    badge_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor("#e2e8f0")),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('LEFTPADDING', (0, 0), (-1, -1), 10),
    ]))
    story.append(badge_table)
    story.append(Spacer(1, 30))

    # Title block
    story.append(Paragraph(title, styles["CoverTitle"]))
    story.append(Paragraph(subtitle, styles["CoverSubtitle"]))

    # Accent decorative line
    accent_bar = Table([[""]], colWidths=[504], rowHeights=[4])
    accent_bar.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#06b6d4")),
    ]))
    story.append(accent_bar)
    story.append(Spacer(1, 40))

    # Metadata card table
    meta_rows = [
        [
            Paragraph("AUDIT TARGET SCOPE", styles["MetaLabel"]),
            Paragraph(f"Internal Network Infrastructure ({total_assets} Active Nodes)", styles["MetaValue"]),
        ],
        [
            Paragraph("SECURITY CLASSIFICATION", styles["MetaLabel"]),
            Paragraph("<font color='#e11d48'><b>STRICTLY CONFIDENTIAL // RESTRICTED</b></font>", styles["MetaValue"]),
        ],
        [
            Paragraph("GENERATED BY", styles["MetaLabel"]),
            Paragraph(f"{_clean_text(user.full_name)} ({_clean_text(user.email)})", styles["MetaValue"]),
        ],
        [
            Paragraph("GENERATION TIMESTAMP", styles["MetaLabel"]),
            Paragraph(f"{datetime.utcnow().strftime('%B %d, %Y - %H:%M:%S UTC')}", styles["MetaValue"]),
        ],
        [
            Paragraph("AUDIT ENGINE VERSION", styles["MetaLabel"]),
            Paragraph("AEGIS SecOps Core v1.0.0 (FastAPI / NVD 2.0 / Nmap)", styles["MetaValue"]),
        ],
    ]
    meta_table = Table(meta_rows, colWidths=[170, 334])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor("#cbd5e1")),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
        ('TOPPADDING', (0, 0), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ('LEFTPADDING', (0, 0), (-1, -1), 12),
        ('RIGHTPADDING', (0, 0), (-1, -1), 12),
    ]))
    story.append(meta_table)

    story.append(Spacer(1, 60))
    notice_text = (
        "<b>Notice of Confidentiality:</b> This document contains sensitive security vulnerability "
        "and threat posture assessment telemetry generated by AEGIS. Unauthorized distribution, copying, "
        "or dissemination is strictly prohibited."
    )
    story.append(Paragraph(notice_text, styles["BodyDark"]))
    story.append(PageBreak())


def _build_executive_pdf(file_path: str, user: User, assets: List[Asset], vulns: List[Vulnerability], date_slug: str):
    """Build Executive Summary PDF."""
    doc = SimpleDocTemplate(
        file_path,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54,
    )
    styles = _get_pdf_stylesheet()
    story = []

    # Cover Page
    _build_cover_page(
        story=story,
        styles=styles,
        title="EXECUTIVE CYBERSECURITY RISK REPORT",
        subtitle="Continuous Vulnerability Intelligence & Organizational Exposure Posture",
        user=user,
        total_assets=len(assets),
        date_slug=date_slug,
    )

    # 1. Executive Summary & Posture Statement
    story.append(Paragraph("1. Executive Risk Posture Overview", styles["SectionHeader"]))

    critical_count = sum(1 for v in vulns if v.severity.value.lower() == "critical")
    high_count = sum(1 for v in vulns if v.severity.value.lower() == "high")
    medium_count = sum(1 for v in vulns if v.severity.value.lower() == "medium")
    low_count = sum(1 for v in vulns if v.severity.value.lower() == "low")
    open_count = sum(1 for v in vulns if v.status.value.lower() == "open")
    mitigated_count = sum(1 for v in vulns if v.status.value.lower() == "mitigated")

    if critical_count > 0:
        posture_statement = (
            f"<b>ELEVATED RISK POSTURE:</b> The AEGIS automated assessment identified <b>{critical_count} Critical</b> "
            f"and <b>{high_count} High</b> severity vulnerabilities requiring immediate remediation. Critical database "
            "and web workloads exhibit unpatched exposures that could allow remote command execution or sensitive "
            "data exfiltration if exposed to hostile networks."
        )
    elif high_count > 0:
        posture_statement = (
            f"<b>HEIGHTENED EXPOSURE:</b> The platform identified <b>{high_count} High</b> severity security findings. "
            "Prompt engineering team intervention is recommended within standard 7-day remediation SLAs."
        )
    elif len(vulns) > 0:
        posture_statement = (
            "<b>MODERATE SECURITY POSTURE:</b> Low-to-moderate severity vulnerabilities detected. Monitored assets are "
            "largely contained, with standard quarterly patch cycles recommended."
        )
    else:
        posture_statement = (
            "<b>SECURE POSTURE:</b> No active vulnerabilities or unpatched network services detected across all monitored infrastructure."
        )

    story.append(Paragraph(posture_statement, styles["BodyDark"]))
    story.append(Spacer(1, 10))

    # Key Metrics Table
    metrics_data = [
        [
            Paragraph("Total Monitored Assets", styles["MetaLabel"]),
            Paragraph(str(len(assets)), styles["TableCellBold"]),
            Paragraph("Total Identified CVEs", styles["MetaLabel"]),
            Paragraph(str(len(vulns)), styles["TableCellBold"]),
        ],
        [
            Paragraph("Critical Severity (CVSS ≥ 9.0)", styles["MetaLabel"]),
            Paragraph(f"<font color='#e11d48'><b>{critical_count}</b></font>", styles["TableCellBold"]),
            Paragraph("High Severity (CVSS 7.0 - 8.9)", styles["MetaLabel"]),
            Paragraph(f"<font color='#ea580c'><b>{high_count}</b></font>", styles["TableCellBold"]),
        ],
        [
            Paragraph("Medium Severity (CVSS 4.0 - 6.9)", styles["MetaLabel"]),
            Paragraph(f"<font color='#ca8a04'><b>{medium_count}</b></font>", styles["TableCellBold"]),
            Paragraph("Low Severity (CVSS < 4.0)", styles["MetaLabel"]),
            Paragraph(f"<font color='#0284c7'><b>{low_count}</b></font>", styles["TableCellBold"]),
        ],
        [
            Paragraph("Open Findings", styles["MetaLabel"]),
            Paragraph(f"<b>{open_count}</b>", styles["TableCellBold"]),
            Paragraph("Mitigated Findings", styles["MetaLabel"]),
            Paragraph(f"<font color='#059669'><b>{mitigated_count}</b></font>", styles["TableCellBold"]),
        ],
    ]
    metrics_table = Table(metrics_data, colWidths=[160, 92, 160, 92])
    metrics_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor("#cbd5e1")),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('LEFTPADDING', (0, 0), (-1, -1), 10),
    ]))
    story.append(metrics_table)
    story.append(Spacer(1, 16))

    # 2. Top Riskiest Vulnerabilities Table
    story.append(Paragraph("2. Top 10 Riskiest Vulnerabilities (Contextual Prioritization)", styles["SectionHeader"]))
    top_vulns = vulns[:10]

    if not top_vulns:
        story.append(Paragraph("<i>No vulnerabilities detected in target scope.</i>", styles["BodyDark"]))
    else:
        table_rows = [[
            Paragraph("CVE IDENTIFIER", styles["TableHead"]),
            Paragraph("SEVERITY", styles["TableHead"]),
            Paragraph("AFFECTED ASSET", styles["TableHead"]),
            Paragraph("BASE CVSS", styles["TableHead"]),
            Paragraph("RISK SCORE", styles["TableHead"]),
            Paragraph("STATUS", styles["TableHead"]),
        ]]
        for v in top_vulns:
            tier = get_risk_tier(v.risk_score)
            table_rows.append([
                Paragraph(f"<b>{_clean_text(v.cve_id)}</b>", styles["TableCellCode"]),
                Paragraph(f"<b>{_clean_text(v.severity.value.upper())}</b>", styles["TableCell"]),
                Paragraph(f"{_clean_text(v.asset.name if v.asset else 'Asset #' + str(v.asset_id))}", styles["TableCell"]),
                Paragraph(f"{v.cvss_score:.1f}", styles["TableCellBold"]),
                Paragraph(f"<b>{v.risk_score:.2f}</b> ({tier})", styles["TableCellBold"]),
                Paragraph(f"{_clean_text(v.status.value.replace('_', ' ').capitalize())}", styles["TableCell"]),
            ])

        vuln_table = Table(table_rows, colWidths=[100, 70, 140, 60, 80, 54])
        t_style = [
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#0f172a")),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor("#cbd5e1")),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
            ('LEFTPADDING', (0, 0), (-1, -1), 6),
            ('RIGHTPADDING', (0, 0), (-1, -1), 6),
        ]
        # Alternate row backgrounds
        for row_idx in range(1, len(table_rows)):
            bg = colors.HexColor("#ffffff") if row_idx % 2 == 1 else colors.HexColor("#f8fafc")
            t_style.append(('BACKGROUND', (0, row_idx), (-1, row_idx), bg))

        vuln_table.setStyle(TableStyle(t_style))
        story.append(vuln_table)

    story.append(Spacer(1, 16))

    # 3. Strategic Recommendations
    story.append(Paragraph("3. Prioritized Strategic Recommendations", styles["SectionHeader"]))
    recs = [
        "<b>Immediate Remediation for Critical Tiers:</b> Patch or isolate vulnerabilities scoring ≥ 9.0 on assets with Criticality 4 or 5 within 24 hours.",
        "<b>Network Hardening & Segmentation:</b> Restrict external network access to sensitive database (port 5432) and backend services to trusted VPN subnets.",
        "<b>Automated Continuous Audits:</b> Schedule daily quick port scans and weekly deep CVE discovery sweeps to prevent configuration drift.",
        "<b>Vulnerability Lifecycle Enforcement:</b> Transition all identified findings from Open to In-Progress with designated system owners in AEGIS.",
    ]
    for r in recs:
        story.append(Paragraph(f"• &nbsp;{r}", styles["BodyDark"]))

    doc.build(story, canvasmaker=NumberedCanvas)


def _build_detailed_pdf(
    file_path: str,
    user: User,
    assets: List[Asset],
    vulns: List[Vulnerability],
    scans: List[Scan],
    date_slug: str,
):
    """Build Comprehensive Technical Audit PDF."""
    doc = SimpleDocTemplate(
        file_path,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54,
    )
    styles = _get_pdf_stylesheet()
    story = []

    # Cover Page
    _build_cover_page(
        story=story,
        styles=styles,
        title="DETAILED TECHNICAL VULNERABILITY AUDIT",
        subtitle="Complete Attack Surface Inventory, Service Fingerprinting & Remediation Playbook",
        user=user,
        total_assets=len(assets),
        date_slug=date_slug,
    )

    # 1. Executive Summary Section
    story.append(Paragraph("1. Assessment Overview & Posture Statement", styles["SectionHeader"]))
    critical_count = sum(1 for v in vulns if v.severity.value.lower() == "critical")
    high_count = sum(1 for v in vulns if v.severity.value.lower() == "high")
    medium_count = sum(1 for v in vulns if v.severity.value.lower() == "medium")
    low_count = sum(1 for v in vulns if v.severity.value.lower() == "low")

    overview_text = (
        f"This comprehensive technical audit details all scanned endpoints across <b>{len(assets)} infrastructure nodes</b>. "
        f"A total of <b>{len(vulns)} vulnerability findings</b> were enriched via NVD API 2.0 and prioritized via the "
        "AEGIS Contextual Risk Engine (CVSS 60% + Asset Criticality 40%)."
    )
    story.append(Paragraph(overview_text, styles["BodyDark"]))
    story.append(Spacer(1, 10))

    # Summary table
    sum_data = [
        [
            Paragraph("Total Monitored Assets", styles["MetaLabel"]),
            Paragraph(str(len(assets)), styles["TableCellBold"]),
            Paragraph("Total CVEs Discovered", styles["MetaLabel"]),
            Paragraph(str(len(vulns)), styles["TableCellBold"]),
        ],
        [
            Paragraph("Critical Vulnerabilities", styles["MetaLabel"]),
            Paragraph(f"<font color='#e11d48'><b>{critical_count}</b></font>", styles["TableCellBold"]),
            Paragraph("High Vulnerabilities", styles["MetaLabel"]),
            Paragraph(f"<font color='#ea580c'><b>{high_count}</b></font>", styles["TableCellBold"]),
        ],
        [
            Paragraph("Medium Vulnerabilities", styles["MetaLabel"]),
            Paragraph(f"<font color='#ca8a04'><b>{medium_count}</b></font>", styles["TableCellBold"]),
            Paragraph("Low Vulnerabilities", styles["MetaLabel"]),
            Paragraph(f"<font color='#0284c7'><b>{low_count}</b></font>", styles["TableCellBold"]),
        ],
    ]
    t = Table(sum_data, colWidths=[160, 92, 160, 92])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor("#cbd5e1")),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
    ]))
    story.append(t)
    story.append(Spacer(1, 14))

    # 2. Asset Inventory Table
    story.append(Paragraph("2. Full Network Asset Inventory", styles["SectionHeader"]))
    asset_rows = [[
        Paragraph("ASSET NAME", styles["TableHead"]),
        Paragraph("IP ADDRESS", styles["TableHead"]),
        Paragraph("TYPE", styles["TableHead"]),
        Paragraph("ENVIRONMENT", styles["TableHead"]),
        Paragraph("CRITICALITY", styles["TableHead"]),
        Paragraph("CVE COUNT", styles["TableHead"]),
    ]]

    # Map asset_id to count of vulns
    vuln_counts = {}
    for v in vulns:
        vuln_counts[v.asset_id] = vuln_counts.get(v.asset_id, 0) + 1

    for a in assets:
        cnt = vuln_counts.get(a.id, 0)
        cnt_str = f"<font color='#e11d48'><b>{cnt}</b></font>" if cnt > 0 else "0"
        asset_rows.append([
            Paragraph(f"<b>{_clean_text(a.name)}</b>", styles["TableCellBold"]),
            Paragraph(f"{_clean_text(a.ip_address)}", styles["TableCellCode"]),
            Paragraph(f"{_clean_text(a.asset_type.value.capitalize())}", styles["TableCell"]),
            Paragraph(f"{_clean_text(a.environment.value.capitalize())}", styles["TableCell"]),
            Paragraph(f"{a.criticality} / 5", styles["TableCellBold"]),
            Paragraph(cnt_str, styles["TableCell"]),
        ])

    assets_table = Table(asset_rows, colWidths=[130, 94, 70, 80, 70, 60])
    assets_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#0f172a")),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor("#cbd5e1")),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
    ]))
    story.append(assets_table)
    story.append(Spacer(1, 14))

    # 3. Scan History Summary Table
    if scans:
        story.append(Paragraph("3. Recent Automated Scan Executions", styles["SectionHeader"]))
        scan_rows = [[
            Paragraph("SCAN ID", styles["TableHead"]),
            Paragraph("TARGET ASSET", styles["TableHead"]),
            Paragraph("TYPE", styles["TableHead"]),
            Paragraph("STATUS", styles["TableHead"]),
            Paragraph("VULNS FOUND", styles["TableHead"]),
            Paragraph("TIMESTAMP", styles["TableHead"]),
        ]]
        for s in scans[:10]:
            asset_label = s.asset.name if s.asset else f"Asset #{s.asset_id}"
            ts = s.started_at.strftime("%Y-%m-%d %H:%M") if s.started_at else "N/A"
            scan_rows.append([
                Paragraph(f"#{s.id}", styles["TableCellCode"]),
                Paragraph(_clean_text(asset_label), styles["TableCell"]),
                Paragraph(_clean_text(s.scan_type.value.upper()), styles["TableCell"]),
                Paragraph(_clean_text(s.status.value.capitalize()), styles["TableCell"]),
                Paragraph(str(s.total_vulns_found), styles["TableCellBold"]),
                Paragraph(ts, styles["TableCell"]),
            ])
        scan_table = Table(scan_rows, colWidths=[54, 140, 60, 70, 80, 100])
        scan_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#0f172a")),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor("#cbd5e1")),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ]))
        story.append(scan_table)
        story.append(Spacer(1, 14))

    # 4. Comprehensive Vulnerability Catalog (Grouped by Asset)
    story.append(Paragraph("4. Comprehensive Vulnerability Catalog", styles["SectionHeader"]))
    story.append(Paragraph(
        "Individual CVE vulnerability dossiers grouped by target host and ordered by composite risk score.",
        styles["BodyDark"]
    ))
    story.append(Spacer(1, 6))

    # Group vulns by asset
    asset_groups = {}
    for v in vulns:
        a_id = v.asset_id
        if a_id not in asset_groups:
            asset_groups[a_id] = []
        asset_groups[a_id].append(v)

    if not vulns:
        story.append(Paragraph("<i>No vulnerabilities detected in target scope.</i>", styles["BodyDark"]))
    else:
        for a_id, a_vulns in asset_groups.items():
            first_v = a_vulns[0]
            a_obj = first_v.asset
            a_name = a_obj.name if a_obj else f"Asset #{a_id}"
            a_ip = a_obj.ip_address if a_obj else "Unknown IP"
            a_crit = a_obj.criticality if a_obj else 3

            # Asset Group Header
            group_banner = [
                [
                    Paragraph(f"<b>TARGET NODE:</b> {_clean_text(a_name)} ({_clean_text(a_ip)})", styles["TableCellBold"]),
                    Paragraph(f"<b>Criticality:</b> {a_crit}/5 &nbsp;|&nbsp; <b>Findings:</b> {len(a_vulns)}", styles["TableCellBold"]),
                ]
            ]
            g_table = Table(group_banner, colWidths=[330, 174])
            g_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#e2e8f0")),
                ('BOX', (0, 0), (-1, -1), 1, colors.HexColor("#94a3b8")),
                ('TOPPADDING', (0, 0), (-1, -1), 5),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
                ('LEFTPADDING', (0, 0), (-1, -1), 8),
            ]))

            vuln_blocks = [g_table, Spacer(1, 6)]

            for v in a_vulns:
                tier = get_risk_tier(v.risk_score)
                svc_desc = f"{_clean_text(v.service or 'Unspecified')} (Port {v.port or 'N/A'})"
                if v.service_version:
                    svc_desc += f" - Version: {_clean_text(v.service_version)}"

                card_rows = [
                    [
                        Paragraph(f"<b>{_clean_text(v.cve_id)}</b>", styles["TableCellCode"]),
                        Paragraph(f"Severity: <b>{v.severity.value.upper()}</b>", styles["TableCell"]),
                        Paragraph(f"CVSS: <b>{v.cvss_score:.1f}</b>", styles["TableCell"]),
                        Paragraph(f"Risk: <b>{v.risk_score:.2f} ({tier})</b>", styles["TableCellBold"]),
                        Paragraph(f"Status: <b>{v.status.value.replace('_', ' ').capitalize()}</b>", styles["TableCell"]),
                    ],
                    [
                        Paragraph("<b>Service / Port:</b>", styles["TableCellBold"]),
                        Paragraph(svc_desc, styles["TableCell"]),
                        Paragraph("", styles["TableCell"]),
                        Paragraph("", styles["TableCell"]),
                        Paragraph("", styles["TableCell"]),
                    ],
                    [
                        Paragraph("<b>Description:</b>", styles["TableCellBold"]),
                        Paragraph(_clean_text(v.description or "No description provided."), styles["TableCell"]),
                        Paragraph("", styles["TableCell"]),
                        Paragraph("", styles["TableCell"]),
                        Paragraph("", styles["TableCell"]),
                    ],
                    [
                        Paragraph("<b>Remediation:</b>", styles["TableCellBold"]),
                        Paragraph(
                            _clean_text(v.remediation or "Review vendor advisories and update to the latest patched software version."),
                            styles["TableCell"]
                        ),
                        Paragraph("", styles["TableCell"]),
                        Paragraph("", styles["TableCell"]),
                        Paragraph("", styles["TableCell"]),
                    ],
                ]

                card_table = Table(card_rows, colWidths=[90, 110, 80, 110, 114])
                card_table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#f1f5f9")),
                    ('SPAN', (1, 1), (-1, 1)),
                    ('SPAN', (1, 2), (-1, 2)),
                    ('SPAN', (1, 3), (-1, 3)),
                    ('BOX', (0, 0), (-1, -1), 0.75, colors.HexColor("#cbd5e1")),
                    ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
                    ('TOPPADDING', (0, 0), (-1, -1), 4),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
                    ('LEFTPADDING', (0, 0), (-1, -1), 6),
                    ('RIGHTPADDING', (0, 0), (-1, -1), 6),
                ]))
                vuln_blocks.append(card_table)
                vuln_blocks.append(Spacer(1, 6))

            story.append(KeepTogether(vuln_blocks))
            story.append(Spacer(1, 10))

    doc.build(story, canvasmaker=NumberedCanvas)


# ==============================================================================
# Excel Compliance Generation
# ==============================================================================

def _build_compliance_excel(
    file_path: str,
    user: User,
    assets: List[Asset],
    vulns: List[Vulnerability],
    date_slug: str,
):
    """Build multi-sheet compliance spreadsheet with auto-filters and formatting."""
    wb = openpyxl.Workbook()

    # Sheet 1: Summary
    ws_summary = wb.active
    ws_summary.title = "Summary"

    # Styling constants
    navy_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    sub_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    bold_white = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_navy = Font(name="Calibri", size=12, bold=True, color="0F172A")
    title_font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="94A3B8")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Title block
    ws_summary.merge_cells("A1:F2")
    ws_summary["A1"] = "AEGIS SEC-OPS | SECURITY COMPLIANCE AUDIT REPORT"
    ws_summary["A1"].font = title_font
    ws_summary["A1"].fill = navy_fill
    ws_summary["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws_summary["A3"] = f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S UTC')}  |  Auditor: {user.full_name} ({user.email})"
    ws_summary["A3"].font = subtitle_font

    # Table 1: Severity Breakdown
    ws_summary["A5"] = "Vulnerability Severity Distribution"
    ws_summary["A5"].font = bold_navy

    ws_summary.append(["Severity Level", "Finding Count", "Share of Total (%)", "Remediation Target SLA"])
    h_row = 6
    for col_idx in range(1, 5):
        cell = ws_summary.cell(row=h_row, column=col_idx)
        cell.font = bold_white
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    sev_counts = {
        "Critical (CVSS ≥ 9.0)": (sum(1 for v in vulns if v.severity.value.lower() == "critical"), "24 Hours"),
        "High (CVSS 7.0 - 8.9)": (sum(1 for v in vulns if v.severity.value.lower() == "high"), "7 Days"),
        "Medium (CVSS 4.0 - 6.9)": (sum(1 for v in vulns if v.severity.value.lower() == "medium"), "30 Days"),
        "Low (CVSS 0.1 - 3.9)": (sum(1 for v in vulns if v.severity.value.lower() == "low"), "90 Days"),
    }
    total_vulns = len(vulns) or 1

    curr_row = 7
    for sev_name, (cnt, sla) in sev_counts.items():
        pct = (cnt / total_vulns) * 100.0
        ws_summary.append([sev_name, cnt, f"{pct:.1f}%", sla])
        for c in range(1, 5):
            ws_summary.cell(row=curr_row, column=c).border = thin_border
        curr_row += 1

    # Table 2: Status Breakdown
    curr_row += 1
    ws_summary.cell(row=curr_row, column=1, value="Remediation Lifecycle Status").font = bold_navy
    curr_row += 1
    ws_summary.cell(row=curr_row, column=1, value="Workflow Status").font = bold_white
    ws_summary.cell(row=curr_row, column=1).fill = header_fill
    ws_summary.cell(row=curr_row, column=1).border = thin_border

    ws_summary.cell(row=curr_row, column=2, value="Findings Count").font = bold_white
    ws_summary.cell(row=curr_row, column=2).fill = header_fill
    ws_summary.cell(row=curr_row, column=2).border = thin_border

    status_counts = {
        "Open": sum(1 for v in vulns if v.status.value.lower() == "open"),
        "In Progress": sum(1 for v in vulns if v.status.value.lower() == "in_progress"),
        "Mitigated": sum(1 for v in vulns if v.status.value.lower() == "mitigated"),
        "False Positive": sum(1 for v in vulns if v.status.value.lower() == "false_positive"),
    }
    curr_row += 1
    for st_name, cnt in status_counts.items():
        ws_summary.append([st_name, cnt])
        for c in range(1, 3):
            ws_summary.cell(row=curr_row, column=c).border = thin_border
        curr_row += 1

    # Adjust summary column widths
    for col in ws_summary.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_summary.column_dimensions[col_letter].width = max(max_len + 4, 16)

    # Sheet 2: Vulnerabilities
    ws_vulns = wb.create_sheet(title="Vulnerabilities")
    vuln_headers = [
        "CVE Identifier",
        "Contextual Risk Score",
        "Risk Tier",
        "Base CVSS",
        "Severity",
        "Status",
        "Affected Asset Name",
        "Asset IP Address",
        "Asset Criticality (1-5)",
        "Service Fingerprint",
        "Port",
        "Service Version",
        "First Discovered",
        "Vulnerability Title / Summary",
        "Remediation Playbook Guidance",
    ]
    ws_vulns.append(vuln_headers)
    ws_vulns.freeze_panes = "A2"

    for col_idx in range(1, len(vuln_headers) + 1):
        cell = ws_vulns.cell(row=1, column=col_idx)
        cell.font = bold_white
        cell.fill = navy_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # Fill styles for severities
    fill_crit = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    fill_high = PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid")
    fill_med = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")
    fill_low = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")

    font_crit = Font(color="991B1B", bold=True)
    font_high = Font(color="9A3412", bold=True)
    font_med = Font(color="854D0E", bold=True)
    font_low = Font(color="075985", bold=True)

    row_idx = 2
    for v in vulns:
        tier = get_risk_tier(v.risk_score)
        a_name = v.asset.name if v.asset else f"Asset #{v.asset_id}"
        a_ip = v.asset.ip_address if v.asset else ""
        a_crit = v.asset.criticality if v.asset else 3
        disc_str = v.first_seen_at.strftime("%Y-%m-%d %H:%M") if v.first_seen_at else ""

        ws_vulns.append([
            v.cve_id,
            float(v.risk_score),
            tier,
            float(v.cvss_score),
            v.severity.value.upper(),
            v.status.value.replace("_", " ").capitalize(),
            a_name,
            a_ip,
            int(a_crit),
            v.service or "Unspecified",
            int(v.port) if v.port else "",
            v.service_version or "",
            disc_str,
            v.title or v.description[:100],
            v.remediation or "Review security advisory.",
        ])

        # Style severity cell (col 5)
        sev_cell = ws_vulns.cell(row=row_idx, column=5)
        sev_val = v.severity.value.lower()
        if sev_val == "critical":
            sev_cell.fill = fill_crit
            sev_cell.font = font_crit
        elif sev_val == "high":
            sev_cell.fill = fill_high
            sev_cell.font = font_high
        elif sev_val == "medium":
            sev_cell.fill = fill_med
            sev_cell.font = font_med
        elif sev_val == "low":
            sev_cell.fill = fill_low
            sev_cell.font = font_low

        for c in range(1, len(vuln_headers) + 1):
            ws_vulns.cell(row=row_idx, column=c).border = thin_border

        row_idx += 1

    # Apply Auto-filter
    ws_vulns.auto_filter.ref = f"A1:{get_column_letter(len(vuln_headers))}{max(row_idx - 1, 1)}"

    for col in ws_vulns.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_vulns.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 50)

    # Sheet 3: Assets
    ws_assets = wb.create_sheet(title="Assets")
    asset_headers = [
        "Asset ID",
        "Asset Name",
        "IP Address",
        "Hostname",
        "Asset Type",
        "Environment",
        "Criticality (1-5)",
        "Owner",
        "Total Vulnerabilities",
        "Critical / High Findings",
    ]
    ws_assets.append(asset_headers)
    ws_assets.freeze_panes = "A2"

    for col_idx in range(1, len(asset_headers) + 1):
        cell = ws_assets.cell(row=1, column=col_idx)
        cell.font = bold_white
        cell.fill = navy_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # Calculate asset vuln counts
    vuln_stats = {}
    for v in vulns:
        if v.asset_id not in vuln_stats:
            vuln_stats[v.asset_id] = {"total": 0, "crit_high": 0}
        vuln_stats[v.asset_id]["total"] += 1
        if v.severity.value.lower() in ["critical", "high"]:
            vuln_stats[v.asset_id]["crit_high"] += 1

    a_row_idx = 2
    for a in assets:
        st = vuln_stats.get(a.id, {"total": 0, "crit_high": 0})
        ws_assets.append([
            a.id,
            a.name,
            a.ip_address,
            a.hostname or "N/A",
            a.asset_type.value.capitalize(),
            a.environment.value.capitalize(),
            a.criticality,
            a.owner or "Security Operations",
            st["total"],
            st["crit_high"],
        ])
        for c in range(1, len(asset_headers) + 1):
            ws_assets.cell(row=a_row_idx, column=c).border = thin_border
        a_row_idx += 1

    ws_assets.auto_filter.ref = f"A1:{get_column_letter(len(asset_headers))}{max(a_row_idx - 1, 1)}"

    for col in ws_assets.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_assets.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 40)

    wb.save(file_path)
